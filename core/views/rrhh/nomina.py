from decimal import Decimal, InvalidOperation
import json
import os
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponse, Http404
from django.db.models import Sum, Count, Q

from core.models import PeriodoNomina, DetalleNominaEmpleado, Empleado


def to_decimal(val, default=Decimal('0')):
    if val is None or val == '':
        return default
    try:
        if isinstance(val, (int, float)):
            return Decimal(str(val))
        clean = str(val).replace('$', '').replace(' ', '').replace(',', '').strip()
        return Decimal(clean)
    except (InvalidOperation, ValueError):
        return default


@login_required
def lista_nominas(request):
    periodos = PeriodoNomina.objects.annotate(
        num_empleados=Count('detalles')
    ).order_by('-ano', '-mes', '-id')

    total_pagado_ano = periodos.filter(estado__in=['APROBADO', 'PAGADO']).aggregate(
        total=Sum('total_neto')
    )['total'] or Decimal('0')

    total_periodos = periodos.count()
    ultimo_periodo = periodos.first()

    context = {
        'periodos': periodos,
        'total_pagado_ano': total_pagado_ano,
        'total_periodos': total_periodos,
        'ultimo_periodo': ultimo_periodo,
        'meses_choices': PeriodoNomina.MESES,
        'ano_actual': datetime.date.today().year,
        'mes_actual': datetime.date.today().month,
    }
    return render(request, 'rrhh/nomina/lista.html', context)


@login_required
def detalle_nomina(request, periodo_id):
    periodo = get_object_or_404(PeriodoNomina, id=periodo_id)
    detalles = periodo.detalles.all().order_by('id')

    empleados_en_nomina = detalles.filter(empleado__isnull=False).values_list('empleado_id', flat=True)
    empleados_disponibles = Empleado.objects.filter(estado='activo').exclude(id__in=empleados_en_nomina).order_by('nombre_completo')

    detalles_data = []
    for d in detalles:
        detalles_data.append({
            'id': d.id,
            'nombre_empleado': d.nombre_empleado,
            'cedula': d.cedula,
            'cargo': d.cargo,
            'salario_basico': float(d.salario_basico),
            'auxilio_transporte_pactado': float(d.auxilio_transporte_pactado),
            'es_aprendiz_sena': d.es_aprendiz_sena,
            'dias_laborados': float(d.dias_laborados),
            'dias_incapacidad': float(d.dias_incapacidad),
            'dias_no_remunerados': float(d.dias_no_remunerados),
            'basico_devengado': float(d.basico_devengado),
            'auxilio_transporte_devengado': float(d.auxilio_transporte_devengado),
            'horas_extras_recargos': float(d.horas_extras_recargos),
            'otros_devengados': float(d.otros_devengados),
            'total_devengado': float(d.total_devengado),
            'deduccion_salud': float(d.deduccion_salud),
            'deduccion_pension': float(d.deduccion_pension),
            'deduccion_afsp': float(d.deduccion_afsp),
            'retencion_fuente': float(d.retencion_fuente),
            'otras_deducciones': float(d.otras_deducciones),
            'total_deducciones': float(d.total_deducciones),
            'neto_pagar': float(d.neto_pagar),
            'observaciones': d.observaciones,
        })

    context = {
        'periodo': periodo,
        'detalles': detalles,
        'empleados_disponibles': empleados_disponibles,
        'detalles_json': json.dumps(detalles_data),
        'smmlv_base': float(periodo.smmlv_base),
        'auxilio_transporte_base': float(periodo.auxilio_transporte_base),
    }
    return render(request, 'rrhh/nomina/detalle.html', context)


@login_required
def generar_periodo_nomina(request):
    if request.method == 'POST':
        ano = int(request.POST.get('ano', datetime.date.today().year))
        mes = int(request.POST.get('mes', datetime.date.today().month))
        tipo = request.POST.get('tipo', 'MENSUAL')
        smmlv = to_decimal(request.POST.get('smmlv_base', '1750905'))
        aux_trans = to_decimal(request.POST.get('auxilio_transporte_base', '249095'))
        observaciones = request.POST.get('observaciones', '').strip()

        try:
            if mes == 12:
                ultimo_dia = datetime.date(ano + 1, 1, 1) - datetime.timedelta(days=1)
            else:
                ultimo_dia = datetime.date(ano, mes + 1, 1) - datetime.timedelta(days=1)
            fecha_inicio = datetime.date(ano, mes, 1)
            fecha_fin = ultimo_dia
        except Exception:
            fecha_inicio = datetime.date.today()
            fecha_fin = datetime.date.today()

        periodo = PeriodoNomina.objects.create(
            ano=ano,
            mes=mes,
            tipo=tipo,
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin,
            smmlv_base=smmlv,
            auxilio_transporte_base=aux_trans,
            observaciones=observaciones,
            estado='BORRADOR'
        )

        empleados_activos = Empleado.objects.filter(estado='activo').order_by('nombre_completo')
        creados = 0
        for emp in empleados_activos:
            salario = to_decimal(emp.salario, default=smmlv)
            cargo_upper = (emp.cargo or '').upper()
            es_aprendiz = 'APRENDIZ' in cargo_upper or 'SENA' in cargo_upper

            detalle = DetalleNominaEmpleado(
                periodo=periodo,
                empleado=emp,
                nombre_empleado=emp.nombre_completo,
                cedula=emp.documento or '',
                cargo=emp.cargo or '',
                salario_basico=salario,
                es_aprendiz_sena=es_aprendiz,
                dias_laborados=Decimal('30'),
                dias_incapacidad=Decimal('0'),
                dias_no_remunerados=Decimal('0'),
                horas_extras_recargos=Decimal('0'),
                otros_devengados=Decimal('0'),
                retencion_fuente=Decimal('0'),
                otras_deducciones=Decimal('0'),
                observaciones=''
            )
            detalle.calcular_automatico()
            detalle.save()
            creados += 1

        periodo.recalcular_totales()
        messages.success(request, f'Nómina de {periodo.get_mes_display()} {periodo.ano} generada con {creados} colaboradores activos.')
        return redirect('detalle_nomina', periodo_id=periodo.id)

    return redirect('lista_nominas')


@login_required
def guardar_nomina_ajax(request, periodo_id):
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': 'Método no permitido'}, status=405)

    periodo = get_object_or_404(PeriodoNomina, id=periodo_id)

    try:
        data = json.loads(request.body)
        filas = data.get('filas', [])

        for item in filas:
            detalle_id = item.get('id')
            detalle = DetalleNominaEmpleado.objects.filter(id=detalle_id, periodo=periodo).first()
            if not detalle:
                continue

            detalle.salario_basico = to_decimal(item.get('salario_basico'))
            detalle.dias_laborados = to_decimal(item.get('dias_laborados', 30))
            detalle.dias_incapacidad = to_decimal(item.get('dias_incapacidad', 0))
            detalle.dias_no_remunerados = to_decimal(item.get('dias_no_remunerados', 0))
            detalle.horas_extras_recargos = to_decimal(item.get('horas_extras_recargos', 0))
            detalle.otros_devengados = to_decimal(item.get('otros_devengados', 0))
            detalle.retencion_fuente = to_decimal(item.get('retencion_fuente', 0))
            detalle.otras_deducciones = to_decimal(item.get('otras_deducciones', 0))
            detalle.observaciones = str(item.get('observaciones', '')).strip()

            detalle.calcular_automatico()
            detalle.save()

        periodo.recalcular_totales()

        return JsonResponse({
            'status': 'ok',
            'message': 'Nómina actualizada exitosamente',
            'totales': {
                'total_devengado': float(periodo.total_devengado),
                'total_deducciones': float(periodo.total_deducciones),
                'total_neto': float(periodo.total_neto),
            }
        })
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=400)


@login_required
def agregar_empleado_nomina(request, periodo_id):
    if request.method == 'POST':
        periodo = get_object_or_404(PeriodoNomina, id=periodo_id)
        empleado_id = request.POST.get('empleado_id')
        nombre_custom = request.POST.get('nombre_custom', '').strip()
        salario = to_decimal(request.POST.get('salario_basico', '0'))
        dias = to_decimal(request.POST.get('dias_laborados', '30'))

        emp = None
        if empleado_id:
            emp = Empleado.objects.filter(id=empleado_id).first()

        if emp:
            nombre = emp.nombre_completo
            cedula = emp.documento or ''
            cargo = emp.cargo or ''
            if salario == Decimal('0'):
                salario = to_decimal(emp.salario, default=periodo.smmlv_base)
        else:
            nombre = nombre_custom
            cedula = request.POST.get('cedula', '').strip()
            cargo = request.POST.get('cargo', '').strip()

        if not nombre:
            messages.error(request, 'Debe indicar el nombre del colaborador.')
            return redirect('detalle_nomina', periodo_id=periodo.id)

        es_aprendiz = 'APRENDIZ' in cargo.upper() or 'SENA' in cargo.upper()
        detalle = DetalleNominaEmpleado(
            periodo=periodo,
            empleado=emp,
            nombre_empleado=nombre,
            cedula=cedula,
            cargo=cargo,
            salario_basico=salario,
            es_aprendiz_sena=es_aprendiz,
            dias_laborados=dias,
            dias_incapacidad=Decimal('0'),
            dias_no_remunerados=Decimal('0'),
            horas_extras_recargos=Decimal('0'),
            otros_devengados=Decimal('0'),
            retencion_fuente=Decimal('0'),
            otras_deducciones=Decimal('0'),
            observaciones=''
        )
        detalle.calcular_automatico()
        detalle.save()

        periodo.recalcular_totales()
        messages.success(request, f'Colaborador {nombre} agregado a la nómina.')
        return redirect('detalle_nomina', periodo_id=periodo.id)

    return redirect('detalle_nomina', periodo_id=periodo_id)


@login_required
def eliminar_empleado_nomina(request, periodo_id, detalle_id):
    if request.method == 'POST':
        periodo = get_object_or_404(PeriodoNomina, id=periodo_id)
        detalle = get_object_or_404(DetalleNominaEmpleado, id=detalle_id, periodo=periodo)
        nombre = detalle.nombre_empleado
        detalle.delete()
        periodo.recalcular_totales()
        messages.success(request, f'Colaborador {nombre} eliminado de la nómina.')
    return redirect('detalle_nomina', periodo_id=periodo_id)


@login_required
def cambiar_estado_nomina(request, periodo_id):
    if request.method == 'POST':
        periodo = get_object_or_404(PeriodoNomina, id=periodo_id)
        nuevo_estado = request.POST.get('estado')
        if nuevo_estado in ['BORRADOR', 'APROBADO', 'PAGADO']:
            periodo.estado = nuevo_estado
            periodo.save(update_fields=['estado'])
            messages.success(request, f'Estado de la nómina actualizado a {periodo.get_estado_display()}.')
    return redirect('detalle_nomina', periodo_id=periodo_id)


@login_required
def importar_nomina_excel(request):
    if request.method == 'POST' and request.FILES.get('archivo_excel'):
        archivo = request.FILES['archivo_excel']
        try:
            wb = openpyxl.load_workbook(archivo, data_only=True)
            hoja_seleccionada = None
            for sname in wb.sheetnames:
                sn_clean = sname.strip().upper()
                if any(m in sn_clean for m in ['SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE', 'ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO', 'AGOSTO']):
                    hoja_seleccionada = wb[sname]
                    nombre_hoja = sname
                    break
            if not hoja_seleccionada:
                hoja_seleccionada = wb.active
                nombre_hoja = hoja_seleccionada.title

            mes = 9
            ano = 2026
            for num_m, nom_m in PeriodoNomina.MESES:
                if nom_m.upper() in nombre_hoja.upper():
                    mes = num_m
                    break
            if '2024' in nombre_hoja:
                ano = 2024
            elif '2025' in nombre_hoja:
                ano = 2025
            elif '2026' in nombre_hoja:
                ano = 2026

            import calendar
            ultimo_dia = calendar.monthrange(ano, mes)[1]
            periodo = PeriodoNomina.objects.create(
                ano=ano,
                mes=mes,
                tipo='MENSUAL',
                fecha_inicio=datetime.date(ano, mes, 1),
                fecha_fin=datetime.date(ano, mes, ultimo_dia),
                smmlv_base=Decimal('1750905'),
                auxilio_transporte_base=Decimal('249095'),
                observaciones=f'Importado de Excel: {archivo.name} ({nombre_hoja})',
                estado='BORRADOR'
            )

            rows = list(hoja_seleccionada.iter_rows(max_col=25, values_only=True))
            header_idx = -1
            for idx, r in enumerate(rows[:10]):
                txt_row = ' '.join([str(x).upper() for x in r if x is not None])
                if 'NOMBRE' in txt_row or 'EMPLEADO' in txt_row:
                    header_idx = idx
                    break

            if header_idx == -1:
                header_idx = 0

            data_rows = rows[header_idx + 1:]
            empleados_db = {e.nombre_completo.strip().upper(): e for e in Empleado.objects.all()}

            creados = 0
            for r in data_rows:
                if not r or not r[0]:
                    continue
                nombre_raw = str(r[0]).strip()
                if not nombre_raw or nombre_raw.upper().startswith('TOTAL') or nombre_raw.upper().startswith('FIRMA'):
                    continue

                basico = to_decimal(r[2] if len(r) > 2 else 0)
                dias = to_decimal(r[4] if len(r) > 4 else 30)
                devengado = to_decimal(r[7] if len(r) > 7 else 0)
                salud = to_decimal(r[8] if len(r) > 8 else 0)
                pension = to_decimal(r[9] if len(r) > 9 else 0)
                afsp = to_decimal(r[11] if len(r) > 11 else 0)
                deducciones = to_decimal(r[12] if len(r) > 12 else 0)
                rte = to_decimal(r[14] if len(r) > 14 else 0)
                neto = to_decimal(r[21] if len(r) > 21 else (r[15] if len(r) > 15 else 0))

                clean_name = nombre_raw.upper().replace('( NO REMUNERADO )', '').replace('( NO REMUNERADA )', '').replace('( INCAPACIDAD )', '').strip()
                emp = None
                for db_name, db_obj in empleados_db.items():
                    if clean_name in db_name or db_name in clean_name:
                        emp = db_obj
                        break

                cedula = emp.documento if emp else (str(r[1]).strip() if len(r) > 1 and r[1] else '')
                cargo = emp.cargo if emp else ''
                es_aprendiz = 'APRENDIZ' in cargo.upper() or 'SENA' in cargo.upper()

                novedad = ''
                if 'NO REMUNERAD' in nombre_raw.upper():
                    novedad = 'Permiso No Remunerado'
                elif 'INCAPACIDAD' in nombre_raw.upper():
                    novedad = 'Incapacidad Médica'

                detalle = DetalleNominaEmpleado(
                    periodo=periodo,
                    empleado=emp,
                    nombre_empleado=nombre_raw,
                    cedula=cedula,
                    cargo=cargo,
                    salario_basico=basico,
                    es_aprendiz_sena=es_aprendiz,
                    dias_laborados=dias,
                    dias_incapacidad=Decimal('0'),
                    dias_no_remunerados=Decimal('0'),
                    basico_devengado=to_decimal(r[5] if len(r) > 5 else 0),
                    auxilio_transporte_devengado=to_decimal(r[6] if len(r) > 6 else 0),
                    total_devengado=devengado,
                    deduccion_salud=salud,
                    deduccion_pension=pension,
                    deduccion_afsp=afsp,
                    retencion_fuente=rte,
                    total_deducciones=deducciones,
                    neto_pagar=neto,
                    observaciones=novedad
                )
                detalle.save()
                creados += 1

            periodo.recalcular_totales()
            messages.success(request, f'Nómina importada exitosamente con {creados} registros desde {archivo.name}.')
            return redirect('detalle_nomina', periodo_id=periodo.id)

        except Exception as e:
            messages.error(request, f'Error al procesar el archivo Excel: {str(e)}')
            return redirect('lista_nominas')

    messages.error(request, 'Debe adjuntar un archivo Excel válido (.xlsx).')
    return redirect('lista_nominas')


@login_required
def exportar_nomina_excel(request, periodo_id):
    periodo = get_object_or_404(PeriodoNomina, id=periodo_id)
    detalles = periodo.detalles.all().order_by('id')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'{periodo.get_mes_display()[:3].upper()} {periodo.ano}'

    font_titulo = Font(name='Arial', size=14, bold=True, color='1E3A8A')
    font_subtitulo = Font(name='Arial', size=10, bold=True, color='475569')
    font_header = Font(name='Arial', size=9, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
    font_data = Font(name='Arial', size=9)
    font_total = Font(name='Arial', size=9, bold=True)
    fill_total = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.merge_cells('A1:L1')
    ws['A1'] = f'COINTECA S.A.S. - NÓMINA GENERAL {periodo.get_mes_display().upper()} DE {periodo.ano}'
    ws['A1'].font = font_titulo
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')

    ws.merge_cells('A2:L2')
    ws['A2'] = f'NIT: 900.563.856-1 | Período: {periodo.fecha_inicio.strftime("%d/%m/%Y")} al {periodo.fecha_fin.strftime("%d/%m/%Y")} | SMMLV: ${periodo.smmlv_base:,.0f} | Aux. Transporte: ${periodo.auxilio_transporte_base:,.0f}'
    ws['A2'].font = font_subtitulo

    headers = [
        'NOMBRE Y APELLIDO', 'CÉDULA', 'CARGO', 'BÁSICO PACTADO', 'DÍAS',
        'BÁSICO DEVENGADO', 'AUX. TRANSPORTE', 'HORAS EXTRAS / OTROS', 'TOTAL DEVENGADO',
        'SALUD (4%)', 'PENSIÓN (4%)', 'FDO. SOLIDARIDAD', 'RETEFUENTE', 'OTRAS DEDUCC.',
        'TOTAL DEDUCCIONES', 'NETO A PAGAR', 'NOVEDADES'
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border_thin

    current_row = 5
    for d in detalles:
        row_vals = [
            d.nombre_empleado,
            d.cedula,
            d.cargo,
            float(d.salario_basico),
            float(d.dias_laborados),
            float(d.basico_devengado),
            float(d.auxilio_transporte_devengado),
            float(d.horas_extras_recargos + d.otros_devengados),
            float(d.total_devengado),
            float(d.deduccion_salud),
            float(d.deduccion_pension),
            float(d.deduccion_afsp),
            float(d.retencion_fuente),
            float(d.otras_deducciones),
            float(d.total_deducciones),
            float(d.neto_pagar),
            d.observaciones
        ]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_thin
            if col_idx in [4, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]:
                cell.number_format = '$#,##0'
                cell.alignment = Alignment(horizontal='right')
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal='center')
        current_row += 1

    total_cell = ws.cell(row=current_row, column=1, value='TOTALES GENERALES')
    total_cell.font = font_total
    total_cell.fill = fill_total
    total_cell.border = border_thin

    for col_idx in range(2, len(headers) + 1):
        c = ws.cell(row=current_row, column=col_idx)
        c.font = font_total
        c.fill = fill_total
        c.border = border_thin
        if col_idx in [6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]:
            col_letter = get_column_letter(col_idx)
            c.value = f'=SUM({col_letter}5:{col_letter}{current_row-1})'
            c.number_format = '$#,##0'
            c.alignment = Alignment(horizontal='right')

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Nomina_Cointeca_{periodo.get_mes_display()}_{periodo.ano}.xlsx"'
    wb.save(response)
    return response


@login_required
def comprobante_pago(request, detalle_id):
    detalle = get_object_or_404(DetalleNominaEmpleado, id=detalle_id)
    periodo = detalle.periodo

    context = {
        'detalle': detalle,
        'periodo': periodo,
        'empresa_nombre': 'COINTECA S.A.S.',
        'empresa_nit': '900.563.856-1',
        'fecha_impresion': datetime.date.today(),
    }
    return render(request, 'rrhh/nomina/comprobante.html', context)
