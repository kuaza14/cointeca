import os
import openpyxl
from django.conf import settings
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError
from django.db.models import Sum
from django.utils import timezone
from decimal import Decimal, InvalidOperation

from core.utils import parse_monto_cop
from core.models import (
    ProyectoFacturacion,
    SeguimientoFacturacion,
    ApoyoManoObra,
    Proyecto,
)

MESES_ORDEN = [
    ('enero', 'Enero'),
    ('febrero', 'Febrero'),
    ('marzo', 'Marzo'),
    ('abril', 'Abril'),
    ('mayo', 'Mayo'),
    ('junio', 'Junio'),
    ('julio', 'Julio'),
    ('agosto', 'Agosto'),
    ('septiembre', 'Septiembre'),
    ('octubre', 'Octubre'),
    ('noviembre', 'Noviembre'),
    ('diciembre', 'Diciembre'),
]

# Vistas de Facturación dentro de Gerencia
@login_required
def facturacion(request):
    proyectos = ProyectoFacturacion.objects.filter(activo=True)

    # Año seleccionado: por defecto el más reciente con datos (o el actual)
    anios_disponibles = list(
        SeguimientoFacturacion.objects
.values_list('anio', flat=True)
.distinct()
.order_by('-anio')
)
    anio_sel = request.GET.get('anio')
    try:
        anio_sel = int(anio_sel) if anio_sel else (
            anios_disponibles[0] if anios_disponibles else timezone.now().year
)
    except (ValueError, TypeError):
        anio_sel = timezone.now().year

    seguimientos = (
        SeguimientoFacturacion.objects
.select_related('proyecto')
.filter(anio=anio_sel)
.order_by('mes')
)

    total_meta = sum([s.meta_facturacion for s in seguimientos]) if seguimientos else 0
    total_real = sum([s.facturacion_real for s in seguimientos]) if seguimientos else 0
    porcentaje_global = round((total_real / total_meta * 100), 1) if total_meta > 0 else 0

    # Consolidado mensual para la matriz anual (Enero a Diciembre)
    matriz_mensual = []
    for mes_cod, mes_nom in MESES_ORDEN:
        segs_mes = [s for s in seguimientos if s.mes.lower() == mes_cod]
        meta_mes = sum([s.meta_facturacion for s in segs_mes])
        real_mes = sum([s.facturacion_real for s in segs_mes])
        pct_mes = round((real_mes / meta_mes * 100), 1) if meta_mes > 0 else 0
        matriz_mensual.append({
            'codigo': mes_cod,
            'nombre': mes_nom,
            'meta': meta_mes,
            'real': real_mes,
            'porcentaje': pct_mes,
            'num_registros': len(segs_mes),
        })

    return render(request, 'gerencia/facturacion/facturacion.html', {
        'proyectos': proyectos,
        'seguimientos': seguimientos,
        'matriz_mensual': matriz_mensual,
        'total_meta': total_meta,
        'total_real': total_real,
        'porcentaje_global': porcentaje_global,
        'anio_seleccionado': anio_sel,
        'anios_disponibles': anios_disponibles,
    })

@login_required
def crear_proyecto_facturacion(request):
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        if nombre:
            ProyectoFacturacion.objects.create(nombre=nombre)
        return redirect('facturacion')
    return render(request, 'gerencia/facturacion/crear_proyecto.html')

@login_required
def registrar_facturacion(request):
    proyectos = ProyectoFacturacion.objects.filter(activo=True)
    if request.method == 'POST':
        try:
            meta = parse_monto_cop(request.POST.get('meta_facturacion'), default=Decimal('0'))
            real = parse_monto_cop(request.POST.get('facturacion_real'), default=Decimal('0'))
        except (InvalidOperation, KeyError):
            messages.error(request, ' Los valores de facturación deben ser numéricos.')
            return render(request, 'gerencia/facturacion/registrar_facturacion.html', {'proyectos': proyectos})

        try:
            SeguimientoFacturacion.objects.create(
                mes=request.POST['mes'],
                anio=request.POST['anio'],
                meta_facturacion=meta,
                facturacion_real=real,
                proyecto_id=request.POST['proyecto']
)
        except IntegrityError:
            messages.error(
                request,
                ' Ya existe un registro de facturación para ese proyecto en ese mes y año.'
)
            return render(request, 'gerencia/facturacion/registrar_facturacion.html', {'proyectos': proyectos})

        return redirect('facturacion')
    return render(request, 'gerencia/facturacion/registrar_facturacion.html', {'proyectos': proyectos})

@login_required
def eliminar_facturacion(request, id):
    seguimiento = get_object_or_404(SeguimientoFacturacion, id=id)
    if request.method == 'POST':
        seguimiento.delete()
        messages.success(request, 'Registro de facturación eliminado.')
    return redirect('facturacion')

@login_required
def editar_facturacion(request, id):
    seguimiento = get_object_or_404(SeguimientoFacturacion, id=id)
    proyectos = ProyectoFacturacion.objects.filter(activo=True)
    if request.method == 'POST':
        try:
            seguimiento.mes = request.POST['mes']
            seguimiento.anio = request.POST['anio']
            seguimiento.meta_facturacion = parse_monto_cop(request.POST.get('meta_facturacion'), default=Decimal('0'))
            seguimiento.facturacion_real = parse_monto_cop(request.POST.get('facturacion_real'), default=Decimal('0'))
            seguimiento.proyecto_id = request.POST['proyecto']
            seguimiento.save()
        except (InvalidOperation, KeyError):
            messages.error(request, ' Los valores de facturación deben ser numéricos.')
        except IntegrityError:
            messages.error(
                request,
                ' Ya existe un registro de facturación para ese proyecto en ese mes y año.'
)
        else:
            return redirect('facturacion')
    return render(request, 'gerencia/facturacion/editar_facturacion.html', {
        'seguimiento': seguimiento,
        'proyectos': proyectos
    })

@login_required
def api_liquidacion_ingenieria(request):
    """
    Retorna el acumulado liquidado en mano de obra desde Ingeniería,
    opcionalmente buscando por proyecto o consolidado general de obras.
    """
    proyecto_fact_id = request.GET.get('proyecto_id')

    total_liquidado = Decimal('0')
    nombre_proyecto = "Todos los proyectos de Ingeniería"

    if proyecto_fact_id:
        p_fact = ProyectoFacturacion.objects.filter(id=proyecto_fact_id).first()
        if p_fact:
            # Match exacto primero; el icontains es el fallback aproximado
            p_ing = (
                Proyecto.objects.filter(nombre__iexact=p_fact.nombre).first()
                or Proyecto.objects.filter(nombre__icontains=p_fact.nombre).first()
)
            if p_ing:
                nombre_proyecto = f"Proyecto: {p_ing.nombre}"
                amos = ApoyoManoObra.objects.filter(apoyo__proyecto=p_ing).select_related('item_mano_obra')
                total_liquidado = sum((a.cantidad * a.item_mano_obra.valor_unitario for a in amos), Decimal('0'))

    if total_liquidado == Decimal('0'):
        # Consolidado general de mano de obra liquidada en apoyos
        amos_todos = ApoyoManoObra.objects.select_related('item_mano_obra').all()
        total_liquidado = sum((a.cantidad * a.item_mano_obra.valor_unitario for a in amos_todos), Decimal('0'))
        if total_liquidado > Decimal('0'):
            nombre_proyecto = "Consolidado Obras de Ingeniería"

    val_entero = int(total_liquidado)
    formateado = f"{val_entero:,}".replace(',', '.')

    return JsonResponse({
        'total': float(total_liquidado),
        'total_formateado': formateado,
        'proyecto_referencia': nombre_proyecto,
    })


@login_required
def exportar_facturacion_excel(request):
    """
    Exporta el libro de seguimiento de facturación (hoja SEGUIMIENTO FACTURACION de GER-FR-02).
    """
    ruta_plantilla = os.path.join(settings.BASE_DIR, 'plantillas_excel', 'formato_tablero_indicadores.xlsx')
    if os.path.exists(ruta_plantilla):
        wb = openpyxl.load_workbook(ruta_plantilla)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = "SEGUIMIENTO FACTURACION"

    ws = wb['SEGUIMIENTO FACTURACION'] if 'SEGUIMIENTO FACTURACION' in wb.sheetnames else wb.active

    seguimientos = SeguimientoFacturacion.objects.select_related('proyecto').all()

    # Mapear cada mes de Enero a Diciembre
    fila_base = 4
    for i, (mes_cod, mes_nom) in enumerate(MESES_ORDEN):
        fila = fila_base + i
        segs_mes = [s for s in seguimientos if s.mes.lower() == mes_cod]
        meta_mes = sum([float(s.meta_facturacion) for s in segs_mes])
        real_mes = sum([float(s.facturacion_real) for s in segs_mes])
        pct_mes = round((real_mes / meta_mes * 100), 1) if meta_mes > 0 else 0

        ws.cell(row=fila, column=2, value=mes_nom)
        if meta_mes > 0:
            ws.cell(row=fila, column=3, value=meta_mes)
            ws.cell(row=fila, column=3).number_format = '$#,##0'
        if real_mes > 0:
            ws.cell(row=fila, column=4, value=real_mes)
            ws.cell(row=fila, column=4).number_format = '$#,##0'
            ws.cell(row=fila, column=5, value=f"{pct_mes}%")

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
    response["Content-Disposition"] = 'attachment; filename="GER-FR-02_Seguimiento_Facturacion.xlsx"'
    wb.save(response)
    return response


