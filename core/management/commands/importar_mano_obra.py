import os
import openpyxl
from decimal import Decimal
from django.core.management.base import BaseCommand
from core.models import ItemManoObra

class Command(BaseCommand):
    help = "Importa el catálogo de 130 actividades de Mano de Obra y sus precios desde proyecciones.xlsx"

    def handle(self, *args, **options):
        excel_path = os.path.join(os.getcwd(), 'proyecciones.xlsx')
        if not os.path.exists(excel_path):
            excel_path = os.path.join(os.getcwd(), 'EJEMPLO EDSON.xlsx')

        if not os.path.exists(excel_path):
            self.stderr.write(self.style.ERROR(f"No se encontró el archivo Excel en {excel_path}"))
            return

        self.stdout.write(f"Leyendo catálogo de Mano de Obra desde {excel_path}...")
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws_resumen = wb['RESUMEN'] if 'RESUMEN' in wb.sheetnames else wb.active
        ws_formulas = wb['9958'] if '9958' in wb.sheetnames else None

        automatic_codes = set()
        if ws_formulas:
            wb_raw = openpyxl.load_workbook(excel_path, data_only=False)
            ws_f_raw = wb_raw['9958']
            for r in range(206, ws_f_raw.max_row + 1):
                code = ws_f_raw.cell(row=r, column=3).value
                formula_cell = ws_f_raw.cell(row=r, column=5).value
                if code and formula_cell and str(formula_cell).startswith('='):
                    automatic_codes.add(str(code).strip())

        mo_items = []
        is_mo = False
        for r in range(1, ws_resumen.max_row + 1):
            desc_val = ws_resumen.cell(row=r, column=2).value
            if desc_val and 'MANO DE OBRA' in str(desc_val).upper():
                is_mo = True
                continue
            if is_mo and desc_val:
                desc = str(desc_val).strip()
                code_val = ws_resumen.cell(row=r, column=3).value
                unit_val = ws_resumen.cell(row=r, column=4).value
                price_val = ws_resumen.cell(row=r, column=5).value

                code = str(code_val).strip() if code_val is not None else f"MO-{r}"
                unit = str(unit_val).strip() if unit_val else "UN"
                try:
                    price = Decimal(str(price_val or 0).strip().replace(',', '.'))
                except Exception:
                    price = Decimal("0")

                mo_items.append({
                    'code': code,
                    'desc': desc,
                    'unit': unit,
                    'price': price
                })

        self.stdout.write(f"Se encontraron {len(mo_items)} actividades de Mano de Obra en la hoja RESUMEN.")

        created_count = 0
        updated_count = 0

        for item in mo_items:
            code = item['code']
            desc = item['desc']
            desc_upper = desc.upper()
            unit = item['unit']
            price = item['price']

            if 'TRANSP' in desc_upper:
                categoria = ItemManoObra.Categorias.TRANSPORTE
            elif any(k in desc_upper for k in ['HINC', 'APLOMADA', 'HUECO', 'BASE DE CONCRETO', 'BASE POST', 'IZAJ', 'VEST']):
                categoria = ItemManoObra.Categorias.HINCADA_APLOMADA
            elif any(k in desc_upper for k in ['DESMON', 'RETIRO']):
                categoria = ItemManoObra.Categorias.DESMONTE
            elif any(k in desc_upper for k in ['LUM', 'PROY', 'TELECELDA', 'REFLECTOR', 'FOTO']):
                categoria = ItemManoObra.Categorias.LUMINARIAS_PROYECTORES
            elif any(k in desc_upper for k in ['CABLE', 'CANALIZ', 'CONDUCTOR', 'ACOMETIDA', 'DUCLO', 'PERF HORIZ']):
                categoria = ItemManoObra.Categorias.REDES_CABLES
            elif any(k in desc_upper for k in ['EXCAV', 'ESCOMBRO', 'CAJA', 'CAMARA', 'CONCRETO', 'ZONA DURA', 'ZONA BL', 'RELL']):
                categoria = ItemManoObra.Categorias.OBRA_CIVIL
            elif any(k in desc_upper for k in ['ABRAZ', 'COLLARIN', 'PERCHA', 'CRUCE', 'DIAGONAL', 'VARILLA', 'TIERRA', 'RETENIDA', 'CORONA', 'REJA', 'TAPA']):
                categoria = ItemManoObra.Categorias.HERRAJES_ACCESORIOS
            else:
                categoria = ItemManoObra.Categorias.OTROS

            tipo_calculo = (
                ItemManoObra.TiposCalculo.AUTOMATICO
                if code in automatic_codes
                else ItemManoObra.TiposCalculo.MANUAL
            )

            obj, created = ItemManoObra.objects.update_or_create(
                codigo=code,
                defaults={
                    'descripcion': desc,
                    'unidad': unit,
                    'valor_unitario': price,
                    'categoria': categoria,
                    'tipo_calculo': tipo_calculo,
                }
            )

            if created:
                created_count += 1
            else:
                updated_count += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"✅ Catálogo de Mano de Obra sincronizado: {created_count} creados, {updated_count} actualizados (Total: {ItemManoObra.objects.count()} actividades)."
            )
        )
