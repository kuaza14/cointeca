import os
import openpyxl
from django.conf import settings
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required

from core.models import (
    IndicadorEstrategico,
    SeguimientoIndicador,
)
from core.helpers.gerencia_seeder import sembrar_indicadores_oficiales

@login_required
def indicadores(request):
    if not IndicadorEstrategico.objects.exists():
        sembrar_indicadores_oficiales()
    items = IndicadorEstrategico.objects.all().order_by('perspectiva')
    return render(request, 'gerencia/indicadores/indicadores.html', {'indicadores': items})

@login_required
def crear_indicador(request):
    if request.method == 'POST':
        IndicadorEstrategico.objects.create(
            perspectiva=request.POST['perspectiva'],
            nombre=request.POST['nombre'],
            definicion=request.POST['definicion'],
            meta_anual=request.POST['meta_anual'],
            frecuencia=request.POST['frecuencia'],
        )
        return redirect('/indicadores/')
    return render(request, 'gerencia/indicadores/crear_indicador.html')

@login_required
def detalle_indicador(request, id):
    indicador = get_object_or_404(IndicadorEstrategico, id=id)
    seguimientos = indicador.seguimientoindicador_set.all().order_by('-fecha')
    return render(request, 'gerencia/indicadores/detalle_indicador.html', {
        'indicador': indicador,
        'seguimientos': seguimientos
    })

@login_required
def agregar_seguimiento(request, id):
    indicador = get_object_or_404(IndicadorEstrategico, id=id)
    if request.method == 'POST':
        SeguimientoIndicador.objects.create(
            indicador=indicador,
            fecha=request.POST['fecha'],
            valor_obtenido=request.POST['valor_obtenido'],
            observaciones=request.POST.get('observaciones', '')
        )
        return redirect(f'/indicadores/{id}/')
    return render(request, 'gerencia/indicadores/agregar_seguimiento.html', {'indicador': indicador})

@login_required
def eliminar_indicador(request, id):
    indicador = get_object_or_404(IndicadorEstrategico, id=id)
    indicador.delete()
    return redirect('/indicadores/')

@login_required
def editar_indicador(request, id):
    indicador = get_object_or_404(IndicadorEstrategico, id=id)

    if request.method == 'POST':
        indicador.perspectiva = request.POST['perspectiva']
        indicador.nombre = request.POST['nombre']
        indicador.definicion = request.POST['definicion']
        indicador.meta_anual = request.POST['meta_anual']
        indicador.frecuencia = request.POST['frecuencia']
        indicador.save()
        return redirect(f'/indicadores/{id}/')

    return render(request, 'gerencia/indicadores/editar_indicador.html', {'indicador': indicador})


@login_required
def eliminar_seguimiento(request, id):
    seguimiento = get_object_or_404(SeguimientoIndicador, id=id)
    indicador_id = seguimiento.indicador.id
    seguimiento.delete()
    return redirect(f'/indicadores/{indicador_id}/')

@login_required
def editar_seguimiento(request, id):    
    seguimiento = get_object_or_404(SeguimientoIndicador, id=id)

    if request.method == 'POST':
        seguimiento.fecha = request.POST['fecha']
        seguimiento.valor_obtenido = request.POST['valor_obtenido']
        seguimiento.observaciones = request.POST.get('observaciones', '')
        seguimiento.save()
        return redirect(f'/indicadores/{seguimiento.indicador.id}/')

    return render(request, 'gerencia/indicadores/editar_seguimiento.html', {'seguimiento': seguimiento})

@login_required
def sembrar_indicadores_view(request):
    sembrar_indicadores_oficiales()
    return redirect('indicadores')

@login_required
def exportar_indicadores_excel(request):
    ruta_plantilla = os.path.join(settings.BASE_DIR, 'plantillas_excel', 'formato_tablero_indicadores.xlsx')
    if os.path.exists(ruta_plantilla):
        wb = openpyxl.load_workbook(ruta_plantilla)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = "INDICADORES ESTRATEGICOS"

    ws = wb[' INDICADORES ESTRATEGICOS'] if ' INDICADORES ESTRATEGICOS' in wb.sheetnames else wb.active

    # Encabezados de auditoría
    ws.cell(row=7, column=7, value="Último Seguimiento")
    ws.cell(row=7, column=8, value="Fecha Seguimiento")

    # Mapear indicadores de la base de datos
    indicadores = IndicadorEstrategico.objects.all().order_by('id')
    fila_base = 9
    for i, ind in enumerate(indicadores):
        fila = fila_base + i
        ws.cell(row=fila, column=2, value=ind.get_perspectiva_display().upper())
        ws.cell(row=fila, column=3, value=ind.nombre)
        ws.cell(row=fila, column=4, value=ind.definicion)
        ws.cell(row=fila, column=5, value=ind.meta_anual)
        ws.cell(row=fila, column=6, value=ind.get_frecuencia_display())

        ultimo_seg = ind.seguimientoindicador_set.order_by('-fecha').first()
        if ultimo_seg:
            ws.cell(row=fila, column=7, value=ultimo_seg.valor_obtenido)
            ws.cell(row=fila, column=8, value=str(ultimo_seg.fecha))

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="GER-FR-02_Tablero_Indicadores_Estrategicos.xlsx"'
    wb.save(response)
    return response

